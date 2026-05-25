from __future__ import annotations

import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "test_corpus" / "public_sample_docs"


DOCS = [
    {
        "kind": "pdf",
        "filename": "pdf_01_public_parks_accessibility.pdf",
        "title": "Public Parks Accessibility Snapshot",
        "subtitle": "Synthetic public test data for local document QA",
        "summary": "This file describes fictional park-accessibility upgrades using non-sensitive public-style data.",
        "records": [
            ["PARK-101", "Harbor View Park", "North District", "Ramp resurfacing", "Completed", "2026-03-18", "$18,400"],
            ["PARK-102", "Cedar Loop Garden", "West District", "Audio wayfinding beacons", "In progress", "2026-06-04", "$24,750"],
            ["PARK-103", "Sunrise Commons", "East District", "Accessible picnic tables", "Scheduled", "2026-08-12", "$9,600"],
        ],
        "questions": [
            ["Which district contains PARK-102?", "West District"],
            ["What upgrade is planned for Sunrise Commons?", "Accessible picnic tables"],
        ],
        "notes": [
            "All site names, budgets, and dates are synthetic.",
            "The status field is intended to test direct table lookup.",
        ],
    },
    {
        "kind": "pdf",
        "filename": "pdf_02_open_library_events.pdf",
        "title": "Open Library Events Calendar",
        "subtitle": "Public sample event records",
        "summary": "A fictional library calendar with event names, rooms, capacities, and public contact desks.",
        "records": [
            ["EVT-210", "Archive Skills Lab", "Maple Room", "2026-04-09", "40 seats", "Reference Desk"],
            ["EVT-211", "Public Data Basics", "Civic Hall", "2026-05-14", "80 seats", "Digital Services"],
            ["EVT-212", "Local History Scan Day", "Oak Room", "2026-07-22", "25 seats", "Collections Desk"],
        ],
        "questions": [
            ["Where is EVT-211 held?", "Civic Hall"],
            ["Which desk owns Local History Scan Day?", "Collections Desk"],
        ],
        "notes": [
            "Calendar entries are invented for repeatable RAG testing.",
            "Use this file for room, date, and capacity questions.",
        ],
    },
    {
        "kind": "pdf",
        "filename": "pdf_03_recycling_pilot_results.pdf",
        "title": "Neighborhood Recycling Pilot Results",
        "subtitle": "Synthetic public operations report",
        "summary": "A short public-style report about fictional recycling pilot zones and measured diversion rates.",
        "records": [
            ["ZONE-A", "Glass sorting carts", "Riverside", "31%", "Operations Team", "High contamination risk"],
            ["ZONE-B", "Food scrap bins", "Hill Market", "44%", "Compost Team", "Best pilot score"],
            ["ZONE-C", "Textile drop boxes", "Union Plaza", "22%", "Reuse Team", "Needs signage refresh"],
        ],
        "questions": [
            ["Which zone had the best pilot score?", "ZONE-B"],
            ["What was the diversion rate for textile drop boxes?", "22%"],
        ],
        "notes": [
            "No real municipal performance data is included.",
            "The diversion rates are synthetic but realistic-looking percentages.",
        ],
    },
    {
        "kind": "pdf",
        "filename": "pdf_04_space_mission_public_facts.pdf",
        "title": "Public Space Mission Fact Cards",
        "subtitle": "Stable public-domain style science facts",
        "summary": "A compact set of widely known public space mission facts for citation and extraction testing.",
        "records": [
            ["MISSION-AP11", "Apollo 11", "Moon landing", "1969", "NASA", "Neil Armstrong and Buzz Aldrin walked on the Moon."],
            ["MISSION-V1", "Voyager 1", "Outer solar system", "1977", "NASA/JPL", "Entered interstellar space after the heliopause crossing."],
            ["MISSION-MSL", "Curiosity Rover", "Mars surface science", "2012", "NASA/JPL", "Landed in Gale Crater."],
        ],
        "questions": [
            ["What year launched Voyager 1?", "1977"],
            ["Where did Curiosity land?", "Gale Crater"],
        ],
        "notes": [
            "Mission records use public, widely known historical facts.",
            "The IDs are added for deterministic QA prompts.",
        ],
    },
    {
        "kind": "pdf",
        "filename": "pdf_05_weather_station_readings.pdf",
        "title": "Open Weather Station Readings",
        "subtitle": "Synthetic public sensor table",
        "summary": "Fictional station readings for testing numeric retrieval, units, and date extraction.",
        "records": [
            ["WX-501", "North Pier", "2026-02-01 09:00", "18.4 C", "62%", "11 km/h"],
            ["WX-502", "Civic Roof", "2026-02-01 09:00", "20.1 C", "48%", "7 km/h"],
            ["WX-503", "Botanic Gate", "2026-02-01 09:00", "17.9 C", "71%", "5 km/h"],
        ],
        "questions": [
            ["Which station had 48% humidity?", "WX-502 / Civic Roof"],
            ["What wind speed was recorded at Botanic Gate?", "5 km/h"],
        ],
        "notes": [
            "Sensor readings are synthetic and should not be used for real weather decisions.",
            "Units are included to test exact-value copying.",
        ],
    },
    {
        "kind": "pdf",
        "filename": "pdf_06_health_workshop_schedule.pdf",
        "title": "Community Health Workshop Schedule",
        "subtitle": "Public sample outreach schedule",
        "summary": "A fictional outreach schedule with public workshop topics, locations, and lead departments.",
        "records": [
            ["HLT-301", "Nutrition Label Reading", "North Clinic", "Wellness Office", "2026-04-16", "Spanish interpretation"],
            ["HLT-302", "Home Safety Basics", "Senior Center", "Injury Prevention", "2026-05-03", "Mobility aids demo"],
            ["HLT-303", "Heat Preparedness", "Civic Library", "Emergency Planning", "2026-06-20", "Cooling-center map"],
        ],
        "questions": [
            ["Which department leads HLT-303?", "Emergency Planning"],
            ["Where is Home Safety Basics scheduled?", "Senior Center"],
        ],
        "notes": [
            "Workshop entries are fictional and public-safe.",
            "Use this document to test department and location extraction.",
        ],
    },
    {
        "kind": "pdf",
        "filename": "pdf_07_transit_maintenance_notices.pdf",
        "title": "Transit Maintenance Notices",
        "subtitle": "Synthetic public service notices",
        "summary": "Fictional transit notices with service windows, route names, and public impacts.",
        "records": [
            ["TRN-410", "Blue Line", "Signal inspection", "2026-04-27 22:00-04:00", "Shuttle buses every 12 minutes", "Track Team"],
            ["TRN-411", "Crosstown Bus", "Stop shelter cleaning", "2026-05-05 09:00-13:00", "No route detour", "Facilities Team"],
            ["TRN-412", "River Tram", "Dock fender replacement", "2026-06-11 06:00-10:00", "Boarding shifted to Pier B", "Marine Team"],
        ],
        "questions": [
            ["What is the public impact for TRN-412?", "Boarding shifted to Pier B"],
            ["Which team owns Blue Line signal inspection?", "Track Team"],
        ],
        "notes": [
            "Transit notices are synthetic and not service advice.",
            "Times are intentionally formatted for exact span retrieval.",
        ],
    },
    {
        "kind": "pdf",
        "filename": "pdf_08_heritage_site_notes.pdf",
        "title": "Open Heritage Site Notes",
        "subtitle": "Public-domain style cultural descriptions",
        "summary": "Short fictional site notes for testing descriptive answers and named locations.",
        "records": [
            ["HRT-701", "Old Canal Lock", "Industrial heritage", "Granite walls and hand-operated gates", "Open-air path"],
            ["HRT-702", "Market Clock Tower", "Civic landmark", "Four-faced clock and bell chamber", "Guided stairs"],
            ["HRT-703", "River Kiln Yard", "Craft history", "Brick kiln bases and clay sorting shed", "Accessible overlook"],
        ],
        "questions": [
            ["What feature is listed for Market Clock Tower?", "Four-faced clock and bell chamber"],
            ["Which site has an accessible overlook?", "River Kiln Yard"],
        ],
        "notes": [
            "These cultural descriptions are synthetic public sample text.",
            "Names are fictional to avoid accidental private references.",
        ],
    },
    {
        "kind": "pdf",
        "filename": "pdf_09_microgrid_operating_log.pdf",
        "title": "Renewable Microgrid Operating Log",
        "subtitle": "Synthetic public infrastructure data",
        "summary": "A fictional microgrid log with generation, storage, and maintenance fields.",
        "records": [
            ["GRID-801", "Solar canopy", "North Depot", "126 kWh", "Battery A", "Normal"],
            ["GRID-802", "Wind kiosk", "Harbor Point", "38 kWh", "Battery B", "Blade inspection due"],
            ["GRID-803", "Library rooftop", "Civic Library", "91 kWh", "Battery C", "Inverter firmware updated"],
        ],
        "questions": [
            ["Which asset generated 91 kWh?", "Library rooftop"],
            ["What note is listed for GRID-802?", "Blade inspection due"],
        ],
        "notes": [
            "Energy readings are made-up public test values.",
            "This document is useful for unit and asset-name QA.",
        ],
    },
    {
        "kind": "pdf",
        "filename": "pdf_10_museum_collection_metadata.pdf",
        "title": "Museum Collection Metadata Extract",
        "subtitle": "Synthetic public catalog records",
        "summary": "Fictional museum catalog entries with accession IDs, materials, and exhibit rooms.",
        "records": [
            ["OBJ-3301", "Harbor Sketchbook", "Paper and graphite", "Gallery 2", "Donated collection", "1890-1905"],
            ["OBJ-3302", "Transit Token Mold", "Steel", "Gallery 4", "Industrial design", "1924"],
            ["OBJ-3303", "Market Banner Fragment", "Cotton textile", "Gallery 1", "Street commerce", "1938"],
        ],
        "questions": [
            ["What material is OBJ-3302 made from?", "Steel"],
            ["Which room displays Harbor Sketchbook?", "Gallery 2"],
        ],
        "notes": [
            "Catalog data is fictional and safe for public sharing.",
            "Accession IDs are deterministic for testing.",
        ],
    },
    {
        "kind": "docx",
        "filename": "docx_01_public_bicycle_counter_report.docx",
        "title": "Public Bicycle Counter Report",
        "subtitle": "Synthetic CC0-style mobility data",
        "summary": "Fictional bicycle counter totals for public RAG testing.",
        "records": [
            ["BIKE-601", "Canal Bridge", "2026-03-01", "1,284 trips", "Morning peak", "Northbound lane"],
            ["BIKE-602", "Museum Avenue", "2026-03-01", "942 trips", "Evening peak", "Protected lane"],
            ["BIKE-603", "Harbor Trail", "2026-03-01", "1,736 trips", "Weekend recreation", "Shared path"],
        ],
        "questions": [
            ["Which bicycle counter recorded 1,736 trips?", "BIKE-603 / Harbor Trail"],
            ["What context is listed for Museum Avenue?", "Evening peak"],
        ],
        "notes": ["Counts are invented for testing.", "No real mobility data is represented."],
    },
    {
        "kind": "docx",
        "filename": "docx_02_public_art_inventory.docx",
        "title": "Public Art Inventory",
        "subtitle": "Synthetic public catalog dataset",
        "summary": "Fictional public art records with artist aliases, materials, and maintenance notes.",
        "records": [
            ["ART-901", "Glass Current", "Avery Stone", "Recycled glass", "Main Plaza", "Clean quarterly"],
            ["ART-902", "Listening Bench", "Mira Vale", "Bronze and oak", "Civic Garden", "Oil wood slats"],
            ["ART-903", "Signal Mural", "Theo Park", "Exterior paint", "Rail Underpass", "UV coating due"],
        ],
        "questions": [
            ["What material is Listening Bench made from?", "Bronze and oak"],
            ["Which artwork has UV coating due?", "Signal Mural"],
        ],
        "notes": ["Artist names are fictional.", "Records are public-safe sample metadata."],
    },
    {
        "kind": "docx",
        "filename": "docx_03_school_garden_observation.docx",
        "title": "School Garden Observation Log",
        "subtitle": "Synthetic public education sample",
        "summary": "A fictional school garden log with plant beds, observations, and actions.",
        "records": [
            ["GDN-110", "Pollinator Bed", "Lavender", "High bee activity", "Add water basin", "2026-04-18"],
            ["GDN-111", "Herb Spiral", "Basil", "Leaf curl spotted", "Check irrigation", "2026-04-19"],
            ["GDN-112", "Rain Garden", "Sedge", "Standing water cleared", "No action", "2026-04-20"],
        ],
        "questions": [
            ["What action is recommended for GDN-111?", "Check irrigation"],
            ["Which bed reported high bee activity?", "Pollinator Bed"],
        ],
        "notes": ["School name and student data are intentionally omitted.", "This file contains no personal data."],
    },
    {
        "kind": "docx",
        "filename": "docx_04_public_wifi_site_list.docx",
        "title": "Public Wi-Fi Site List",
        "subtitle": "Synthetic connectivity dataset",
        "summary": "Fictional public Wi-Fi access points, bandwidth classes, and service notes.",
        "records": [
            ["WIFI-220", "Civic Library", "Indoor", "High", "7 AM-9 PM", "Content filter active"],
            ["WIFI-221", "Market Square", "Outdoor", "Medium", "24 hours", "Weatherproof cabinet"],
            ["WIFI-222", "Transit Center", "Indoor", "High", "5 AM-midnight", "Captive portal refreshed"],
        ],
        "questions": [
            ["Which Wi-Fi site is available 24 hours?", "Market Square"],
            ["What note is listed for WIFI-222?", "Captive portal refreshed"],
        ],
        "notes": ["Network records are fictional.", "No credentials, keys, or real network details are included."],
    },
    {
        "kind": "docx",
        "filename": "docx_05_public_tree_canopy_survey.docx",
        "title": "Public Tree Canopy Survey",
        "subtitle": "Synthetic environmental sample",
        "summary": "Fictional street-tree observations with species, canopy class, and maintenance recommendations.",
        "records": [
            ["TREE-704", "Elm", "Broadleaf", "Good", "Main Street", "Mulch ring refresh"],
            ["TREE-705", "Ginkgo", "Medium canopy", "Fair", "Station Road", "Prune low branch"],
            ["TREE-706", "London plane", "Large canopy", "Excellent", "River Walk", "Monitor roots"],
        ],
        "questions": [
            ["Which tree needs pruning?", "TREE-705 / Ginkgo"],
            ["What condition is listed for London plane?", "Excellent"],
        ],
        "notes": ["Tree IDs and observations are synthetic.", "The data is safe to share publicly."],
    },
    {
        "kind": "docx",
        "filename": "docx_06_open_budget_training_examples.docx",
        "title": "Open Budget Training Examples",
        "subtitle": "Synthetic civic finance practice data",
        "summary": "Fictional budget rows designed for public training and exact amount retrieval.",
        "records": [
            ["BUD-310", "Library outreach", "Education", "$42,000", "Approved", "Quarter 2"],
            ["BUD-311", "Park benches", "Public works", "$18,500", "Draft", "Quarter 3"],
            ["BUD-312", "Heat response supplies", "Emergency planning", "$27,250", "Approved", "Quarter 2"],
        ],
        "questions": [
            ["What amount is assigned to Heat response supplies?", "$27,250"],
            ["Which budget item is still draft?", "Park benches"],
        ],
        "notes": ["Amounts are synthetic and not financial advice.", "Use for public RAG testing only."],
    },
    {
        "kind": "docx",
        "filename": "docx_07_public_market_vendor_directory.docx",
        "title": "Public Market Vendor Directory",
        "subtitle": "Synthetic public directory",
        "summary": "Fictional vendor records with stall IDs, product categories, and operating days.",
        "records": [
            ["MRK-501", "Orchard Table", "Fruit", "Stall A4", "Saturday", "Cold storage requested"],
            ["MRK-502", "Clay Cup Studio", "Ceramics", "Stall B2", "Sunday", "Fragile goods signage"],
            ["MRK-503", "River Greens", "Vegetables", "Stall C1", "Saturday and Sunday", "Compost pickup"],
        ],
        "questions": [
            ["Which vendor occupies Stall B2?", "Clay Cup Studio"],
            ["What note is listed for River Greens?", "Compost pickup"],
        ],
        "notes": ["Vendor names are fictional.", "No real business data is included."],
    },
    {
        "kind": "docx",
        "filename": "docx_08_open_grant_award_summaries.docx",
        "title": "Open Grant Award Summaries",
        "subtitle": "Synthetic public grants dataset",
        "summary": "Fictional grant award summaries with program names, award amounts, and public outcomes.",
        "records": [
            ["GRANT-720", "Neighborhood Shade", "Urban forestry", "$75,000", "Plant 220 trees", "Awarded"],
            ["GRANT-721", "Digital Navigator", "Library access", "$58,000", "Train 400 residents", "Awarded"],
            ["GRANT-722", "Pocket Park Repair", "Public works", "$36,500", "Repair 12 benches", "Pending"],
        ],
        "questions": [
            ["Which grant plans to train 400 residents?", "Digital Navigator"],
            ["What is the status of Pocket Park Repair?", "Pending"],
        ],
        "notes": ["Grant names and amounts are synthetic.", "The format mimics public grant summaries."],
    },
    {
        "kind": "docx",
        "filename": "docx_09_public_safety_drill_calendar.docx",
        "title": "Public Safety Drill Calendar",
        "subtitle": "Synthetic preparedness schedule",
        "summary": "Fictional drill calendar records with agencies, dates, and exercise objectives.",
        "records": [
            ["DRILL-140", "Flood tabletop", "Emergency Planning", "2026-05-17", "Test alert routing", "Civic Hall"],
            ["DRILL-141", "Shelter setup", "Community Services", "2026-06-08", "Assemble 60 cots", "Gym Annex"],
            ["DRILL-142", "Radio check", "Volunteer Corps", "2026-06-21", "Confirm repeater coverage", "Hill Station"],
        ],
        "questions": [
            ["Where is the shelter setup drill?", "Gym Annex"],
            ["Which drill tests alert routing?", "Flood tabletop"],
        ],
        "notes": ["Preparedness entries are fictional and non-operational.", "No real emergency procedures are included."],
    },
    {
        "kind": "docx",
        "filename": "docx_10_public_water_quality_summary.docx",
        "title": "Public Water Quality Summary",
        "subtitle": "Synthetic monitoring records",
        "summary": "Fictional water quality observations with public-style station IDs and values.",
        "records": [
            ["WQ-880", "North Creek", "pH 7.2", "Turbidity 3.1 NTU", "2026-03-10", "Normal"],
            ["WQ-881", "Lake Outlet", "pH 7.8", "Turbidity 5.4 NTU", "2026-03-10", "Watch"],
            ["WQ-882", "South Inlet", "pH 6.9", "Turbidity 2.7 NTU", "2026-03-10", "Normal"],
        ],
        "questions": [
            ["Which station has Watch status?", "WQ-881 / Lake Outlet"],
            ["What turbidity is listed for South Inlet?", "2.7 NTU"],
        ],
        "notes": ["Values are synthetic and not regulatory measurements.", "Use only for QA testing."],
    },
]


def ensure_output_dir() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)


def public_notice() -> str:
    return (
        "Source status: Synthetic public test data. License note: CC0-style sample content "
        "created for local RAG testing. No private, personal, confidential, or operationally "
        "valid data is included."
    )


def make_pdf(spec: dict) -> None:
    path = OUT_DIR / spec["filename"]
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "Title",
        parent=styles["Title"],
        textColor=colors.HexColor("#06343A"),
        fontName="Helvetica-Bold",
        fontSize=22,
        leading=26,
        spaceAfter=10,
    )
    subtitle = ParagraphStyle(
        "Subtitle",
        parent=styles["BodyText"],
        textColor=colors.HexColor("#587078"),
        fontSize=10,
        leading=14,
        spaceAfter=14,
    )
    body = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontSize=10.5,
        leading=14,
        spaceAfter=8,
    )
    heading = ParagraphStyle(
        "Heading",
        parent=styles["Heading2"],
        textColor=colors.HexColor("#008B7A"),
        fontSize=14,
        leading=18,
        spaceBefore=12,
        spaceAfter=8,
    )

    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=0.7 * inch,
        rightMargin=0.7 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
        title=spec["title"],
        author="VaultIQ Public Test Corpus",
    )
    story = [
        Paragraph(spec["title"], title),
        Paragraph(spec["subtitle"], subtitle),
        Paragraph(public_notice(), body),
        Paragraph(spec["summary"], body),
        Paragraph("Records", heading),
    ]

    header = ["ID", "Name", "Field A", "Field B", "Field C", "Field D", "Field E"]
    rows = [header] + spec["records"]
    table = Table(rows, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAFBF5")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#06343A")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.2),
                ("LEADING", (0, 0), (-1, -1), 10),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#BFDAD4")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FAFDFB")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.extend([table, Spacer(1, 0.12 * inch), Paragraph("Good test questions", heading)])
    for question, answer in spec["questions"]:
        story.append(Paragraph(f"<b>Q:</b> {question}<br/><b>Expected:</b> {answer}", body))
    story.append(Paragraph("Notes", heading))
    for note in spec["notes"]:
        story.append(Paragraph(f"- {note}", body))
    doc.build(story)


def make_docx(spec: dict) -> None:
    path = OUT_DIR / spec["filename"]
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(15)
    styles["Heading 1"].font.bold = True

    header = section.header.paragraphs[0]
    header.text = "VaultIQ public test corpus"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    title = doc.add_paragraph(style="Title")
    title.add_run(spec["title"])
    sub = doc.add_paragraph()
    sub.add_run(spec["subtitle"]).italic = True
    doc.add_paragraph(public_notice())
    doc.add_paragraph(spec["summary"])

    doc.add_heading("Records", level=1)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["ID", "Name", "Field A", "Field B", "Field C", "Field D", "Field E"]
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row in spec["records"]:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value

    doc.add_heading("Good test questions", level=1)
    for question, answer in spec["questions"]:
        p = doc.add_paragraph()
        p.add_run("Q: ").bold = True
        p.add_run(question)
        p.add_run("\nExpected: ").bold = True
        p.add_run(answer)

    doc.add_heading("Notes", level=1)
    for note in spec["notes"]:
        doc.add_paragraph(note, style="List Bullet")

    doc.core_properties.title = spec["title"]
    doc.core_properties.subject = "Synthetic public test data for VaultIQ"
    doc.core_properties.author = "VaultIQ Public Test Corpus"
    doc.save(path)


def write_manifest() -> None:
    rows = []
    for spec in DOCS:
        for question, answer in spec["questions"]:
            rows.append(
                {
                    "filename": spec["filename"],
                    "title": spec["title"],
                    "question": question,
                    "expected_answer": answer,
                    "data_status": "synthetic public test data",
                }
            )

    with (OUT_DIR / "answer_key.csv").open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)

    with (OUT_DIR / "manifest.json").open("w", encoding="utf-8") as f:
        json.dump(
            {
                "name": "VaultIQ public sample document corpus",
                "license_note": public_notice(),
                "file_count": len(DOCS),
                "pdf_count": sum(1 for item in DOCS if item["kind"] == "pdf"),
                "docx_count": sum(1 for item in DOCS if item["kind"] == "docx"),
                "files": [
                    {
                        "filename": item["filename"],
                        "title": item["title"],
                        "kind": item["kind"],
                        "questions": item["questions"],
                    }
                    for item in DOCS
                ],
            },
            f,
            indent=2,
        )

    readme = OUT_DIR / "README.txt"
    readme.write_text(
        "VaultIQ public sample document corpus\n"
        "====================================\n\n"
        f"{public_notice()}\n\n"
        "Import this folder into VaultIQ to test PDF/DOCX extraction, chunking, retrieval, and answer quality.\n"
        "Use answer_key.csv for expected answers.\n",
        encoding="utf-8",
    )


def main() -> None:
    ensure_output_dir()
    for existing in OUT_DIR.glob("*"):
        if existing.is_file():
            existing.unlink()
    for spec in DOCS:
        if spec["kind"] == "pdf":
            make_pdf(spec)
        else:
            make_docx(spec)
    write_manifest()
    print(f"Wrote {len(DOCS)} documents to {OUT_DIR}")


if __name__ == "__main__":
    main()
